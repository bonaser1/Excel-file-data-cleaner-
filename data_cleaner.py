import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import os

def clean_excel_file(input_path, output_path):
    print(f"\n{'='*50}")
    print(f"Cleaning: {os.path.basename(input_path)}")
    print(f"{'='*50}")

    df = pd.read_excel(input_path)
    original_rows = len(df)
    print(f"Original rows: {original_rows}")

    # 1. Clean whitespace from string columns
    str_cols = df.select_dtypes(include='object').columns
    for col in str_cols:
        df[col] = df[col].astype(str).str.strip()
        df[col] = df[col].replace('nan', '')
        df[col] = df[col].replace('None', '')

    # 2. Standardize text (title case for name/city/dept columns)
    for col in str_cols:
        col_lower = col.lower()
        if any(word in col_lower for word in ['name', 'city', 'department', 'dept', 'region', 'status', 'product']):
            df[col] = df[col].str.title()

    # 3. Remove duplicate rows
    df_before_dedup = len(df)
    df = df.drop_duplicates()
    duplicates_removed = df_before_dedup - len(df)
    print(f"Duplicates removed: {duplicates_removed}")

    # 4. Remove rows where ALL values are empty
    df = df.replace('', pd.NA)
    df = df.dropna(how='all')

    # 5. Reset index
    df = df.reset_index(drop=True)
    final_rows = len(df)
    print(f"Final rows: {final_rows}")
    print(f"Total rows cleaned: {original_rows - final_rows}")

    # ===== Write clean file with formatting =====
    wb = openpyxl.Workbook()

    # Sheet 1: Clean Data
    ws_data = wb.active
    ws_data.title = "Clean Data"

    # Header style
    header_fill = PatternFill("solid", start_color="1F4E79", end_color="1F4E79")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_align = Alignment(horizontal="center", vertical="center")

    # Write headers
    for col_idx, col_name in enumerate(df.columns, 1):
        cell = ws_data.cell(row=1, column=col_idx, value=col_name)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_align

    # Write data
    row_fill_1 = PatternFill("solid", start_color="EBF3FB", end_color="EBF3FB")
    row_fill_2 = PatternFill("solid", start_color="FFFFFF", end_color="FFFFFF")
    data_font = Font(size=10)

    for row_idx, row in df.iterrows():
        for col_idx, value in enumerate(row, 1):
            cell = ws_data.cell(row=row_idx + 2, column=col_idx, value=value if pd.notna(value) else "")
            cell.fill = row_fill_1 if row_idx % 2 == 0 else row_fill_2
            cell.font = data_font
            cell.alignment = Alignment(vertical="center")

    # Auto column width
    for col_idx, col_name in enumerate(df.columns, 1):
        max_len = max(len(str(col_name)), df[col_name].astype(str).str.len().max())
        ws_data.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 4, 30)

    ws_data.row_dimensions[1].height = 25
    ws_data.freeze_panes = "A2"

    # Sheet 2: Summary Report
    ws_summary = wb.create_sheet("Cleaning Report")
    ws_summary.column_dimensions['A'].width = 30
    ws_summary.column_dimensions['B'].width = 20

    title_font = Font(bold=True, size=14, color="1F4E79")
    ws_summary['A1'] = "Data Cleaning Summary Report"
    ws_summary['A1'].font = title_font

    summary_data = [
        ("", ""),
        ("File", os.path.basename(input_path)),
        ("Original Rows", original_rows),
        ("Duplicates Removed", duplicates_removed),
        ("Final Clean Rows", final_rows),
        ("Rows Cleaned", original_rows - final_rows),
        ("Columns", len(df.columns)),
    ]

    label_font = Font(bold=True, size=11)
    value_font = Font(size=11)

    for row_idx, (label, value) in enumerate(summary_data, 2):
        ws_summary.cell(row=row_idx, column=1, value=label).font = label_font
        ws_summary.cell(row=row_idx, column=2, value=value).font = value_font

    wb.save(output_path)
    print(f"Saved: {output_path}")
    return final_rows


if __name__ == "__main__":
    os.makedirs("/cleaned_output", exist_ok=True)

    files = [
        ("dirty_sales_data.xlsx", "clean_sales_data.xlsx")
    ]

    for input_file, output_file in files:
        clean_excel_file(input_file, output_file)

    print("\nAll files cleaned successfully! ✅")